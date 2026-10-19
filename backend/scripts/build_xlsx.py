"""Build data/questoes_geradas.xlsx: the fully auditable 1,788-item corpus.

For every item it holds the generated content, the curricular context, each structural
criterion (C1-C5), the NLP metrics as stored and as recomputed from scratch, duplicate
flags, both LLM judges (principal: glm-4-plus, with raw response, tokens and timestamp;
secondary: glm-4.5, the generator itself) and the exact generation and judge prompts,
rebuilt with the same functions the pipeline used. Further sheets carry the Reference
Matrix, per-skill aggregates, every statistical result and a data dictionary.

Sources (all in data/): questoes_export.json, julgamento_glm-4-plus.jsonl,
matriz_referencia.json, conferencia_extracao_matriz.json and the resultado_*.json files.

Usage (from backend/):
    uv run --with openpyxl python -m scripts.build_xlsx
"""

import ast
import difflib
import json
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

backend_dir = Path(__file__).parent.parent
sys.path.append(str(backend_dir))

from src.modules.avaliacao import nlp  # noqa: E402
from src.modules.avaliacao.judge_prompt import build_judge_prompt  # noqa: E402
from src.modules.geracao.prompts import build_context_block, build_prompt  # noqa: E402
from src.modules.geracao.schemas import CondicaoExperimental  # noqa: E402
from src.modules.matriz_referencia.schemas import HabilidadeContexto, ObjetoConhecimentoRead  # noqa: E402
from src.modules.validacao.schemas import ItemGerado  # noqa: E402
from src.modules.validacao.service import ValidacaoService  # noqa: E402

DATA_DIR = backend_dir / "data"
CONDICOES = ["com_grafo", "com_fewshot_sem_grafo", "sem_grafo"]
REPO = "https://github.com/Jeova-1704/GRAPH-DRIVEN-EXAM-GEN"
AZUL = PatternFill("solid", start_color="1F4E78")


def ler_json(nome: str):
    return json.loads((DATA_DIR / nome).read_text(encoding="utf-8"))


registros = ler_json("questoes_export.json")
matriz = ler_json("matriz_referencia.json")
juiz_principal: dict[str, dict] = {}
with (DATA_DIR / "julgamento_glm-4-plus.jsonl").open(encoding="utf-8") as f:
    for linha in f:
        if linha.strip():
            j = json.loads(linha)
            juiz_principal[j["id"]] = j

# --- Curricular context per habilidade, exactly as the pipeline builds it ---
habilidades: dict[str, HabilidadeContexto] = {}
tokens_oc: dict[str, list[str]] = {}
for area in matriz["areas"]:
    ocs = {oc["codigo"]: ObjetoConhecimentoRead(**oc) for oc in area["objetos_conhecimento"]}
    for comp in area["competencias"]:
        objetos = [ocs[c] for c in comp["objetos_conhecimento"]]
        texto_oc = " ".join(f"{o.titulo} {o.descricao}" for o in objetos)
        for hab in comp["habilidades"]:
            habilidades[hab["codigo"]] = HabilidadeContexto(
                habilidade_codigo=hab["codigo"],
                habilidade_descricao=hab["descricao"],
                competencia_numero=comp["numero"],
                competencia_descricao=comp["descricao"],
                area_codigo=area["codigo"],
                area_nome=area["nome"],
                objetos_conhecimento=objetos,
            )
            tokens_oc[hab["codigo"]] = nlp.tokenizar(texto_oc)


def item_de(q: dict) -> ItemGerado:
    return ItemGerado(
        enunciado=q["enunciado"],
        alternativas={letra: q[f"alternativa_{letra.lower()}"] for letra in "ABCDE"},
        gabarito=q["gabarito"],
        justificativa=q["justificativa"] or "",
    )


def numero(v):
    """Values from the JSON export may come as strings; store them as numbers in the sheet."""
    if v is None or isinstance(v, bool):
        return v
    try:
        f = float(v)
    except (TypeError, ValueError):
        return v
    return int(f) if f.is_integer() and not isinstance(v, float) else f


def media(valores: list) -> float | None:
    vals = [numero(v) for v in valores]
    return sum(vals) / len(vals) if vals and all(v is not None for v in vals) else None


# --- Duplicates within each cell (habilidade x condicao) ---
por_celula: dict[tuple[str, str], list[dict]] = defaultdict(list)
for r in registros:
    por_celula[(r["habilidade_codigo"], r["origem"])].append(r)
primeira_ocorrencia: dict[str, str | None] = {}
maior_similaridade: dict[str, float] = {}
for itens in por_celula.values():
    vistos: dict[str, str] = {}
    for q in itens:
        chave = q["enunciado"].strip()
        primeira_ocorrencia[q["id"]] = vistos.get(chave)
        vistos.setdefault(chave, q["id"])
        outras = [o for o in itens if o["id"] != q["id"]]
        maior_similaridade[q["id"]] = max(
            (difflib.SequenceMatcher(None, q["enunciado"], o["enunciado"]).ratio() for o in outras), default=0.0
        )

validador = ValidacaoService()
CRITERIOS = [
    ("C1_cinco_alternativas", "C1 - cinco alternativas A-E"),
    ("C2_gabarito_valido", "C2 - gabarito valido"),
    ("C3_sem_duplicatas", "C3 - sem alternativas duplicadas"),
    ("C4_enunciado_minimo", "C4 - enunciado >= 50 caracteres"),
    ("C5_justificativa_minima", "C5 - justificativa >= 100 caracteres"),
]
DIMS_SECUNDARIO = [
    "judge_d1_alinhamento_habilidade",
    "judge_d2_adequacao_objetos_conhecimento",
    "judge_d3_plausibilidade_distratores",
    "judge_d4_nivel_escolaridade",
    "judge_d5_aderencia_formato_enem",
]

# (chave, rotulo, largura, descricao para o dicionario de dados)
colunas = [
    ("id", "ID", 24, "Identificador unico (UUID) do no :Questao no grafo Neo4j."),
    ("origem", "Condicao", 20, "Condicao experimental: com_grafo, com_fewshot_sem_grafo ou sem_grafo."),
    ("area_codigo", "Area (codigo)", 9, "Area do conhecimento da habilidade-alvo (LC, MT, CN, CH)."),
    ("area_nome", "Area (nome)", 22, "Nome da area do conhecimento."),
    ("competencia_numero", "Competencia", 10, "Numero da competencia de area da habilidade-alvo."),
    ("competencia_descricao", "Competencia (descricao)", 40, "Descricao oficial da competencia (Matriz do INEP)."),
    ("habilidade_codigo", "Habilidade", 10, "Codigo da habilidade-alvo (aresta :AVALIA)."),
    ("habilidade_descricao", "Habilidade (descricao)", 40, "Descricao oficial da habilidade-alvo (Matriz do INEP)."),
    (
        "ocs",
        "Objetos de conhecimento da competencia",
        40,
        "Titulos dos objetos de conhecimento ligados a competencia; base de O(h).",
    ),
    ("enunciado", "Enunciado", 60, "Enunciado gerado pelo LLM (glm-4.5)."),
    ("alternativa_a", "Alternativa A", 30, "Alternativa A gerada."),
    ("alternativa_b", "Alternativa B", 30, "Alternativa B gerada."),
    ("alternativa_c", "Alternativa C", 30, "Alternativa C gerada."),
    ("alternativa_d", "Alternativa D", 30, "Alternativa D gerada."),
    ("alternativa_e", "Alternativa E", 30, "Alternativa E gerada."),
    ("gabarito", "Gabarito", 8, "Alternativa correta indicada pelo gerador."),
    ("justificativa", "Justificativa (gerador)", 40, "Justificativa do gabarito escrita pelo gerador."),
    (
        "validacao_estrutural_ok",
        "Validacao estrutural OK (registrada)",
        12,
        "Resultado da validacao gravado no grafo no momento da geracao.",
    ),
    *[
        (chave, rotulo, 11, f"Criterio {rotulo} reavaliado agora com o mesmo codigo (ValidacaoService).")
        for chave, rotulo in CRITERIOS
    ],
    (
        "dup",
        "Enunciado repetido na celula",
        11,
        "Sim se o enunciado repete literalmente o de outro item da mesma habilidade e condicao.",
    ),
    ("dup_de", "Repete o item (ID)", 24, "ID da primeira ocorrencia do mesmo enunciado na celula."),
    (
        "sim_max",
        "Maior similaridade com outro item da celula",
        12,
        "Similaridade textual maxima (difflib, 0 a 1) com outro enunciado da mesma celula.",
    ),
    ("cobertura_lexical", "Cobertura lexical C(i,h)", 13, "C(i,h) gravado no grafo."),
    (
        "c_recalc",
        "C(i,h) recalculado",
        13,
        "C(i,h) recalculado agora a partir do texto do item e da Matriz (nlp.cobertura_lexical).",
    ),
    (
        "especificidade_curricular",
        "Especificidade curricular E(i,h)",
        13,
        "E(i,h) gravado no grafo (media das outras 119 habilidades com cobertura positiva).",
    ),
    ("e_recalc", "E(i,h) recalculado", 13, "E(i,h) recalculado agora (nlp.especificidade_curricular)."),
    ("ttr", "TTR", 9, "Diversidade lexical (type-token ratio)."),
    ("fkgl", "FKGL", 9, "Legibilidade Flesch-Kincaid (calibrada para o ingles; so comparacao relativa)."),
    ("p_modelo", "Juiz principal - modelo", 12, "Modelo do juiz principal (distinto do gerador)."),
    ("p_julgado_em", "Juiz principal - data/hora (UTC)", 20, "Momento do julgamento."),
    ("p_temperatura", "Juiz principal - temperatura", 10, "Temperatura da chamada."),
    ("p_max_tokens", "Juiz principal - max_tokens", 10, "Limite de tokens da resposta."),
    ("p_media", "Juiz principal - media", 10, "Media simples de D1 a D5."),
    ("p_d1", "Principal D1 - Alinhamento a habilidade", 11, "Nota 1-5."),
    ("p_d2", "Principal D2 - Adequacao dos OCs", 11, "Nota 1-5."),
    ("p_d3", "Principal D3 - Plausibilidade dos distratores", 11, "Nota 1-5."),
    ("p_d4", "Principal D4 - Nivel de escolaridade", 11, "Nota 1-5."),
    ("p_d5", "Principal D5 - Formato ENEM", 11, "Nota 1-5."),
    ("p_justificativa", "Principal - justificativa", 40, "Justificativa textual do juiz principal."),
    ("p_tok_prompt", "Principal - tokens de entrada", 10, "Tokens do prompt (informado pela API)."),
    ("p_tok_resp", "Principal - tokens de saida", 10, "Tokens da resposta (informado pela API)."),
    ("p_tok_total", "Principal - tokens total", 10, "Total de tokens."),
    ("p_bruta", "Principal - resposta bruta", 40, "Texto integral devolvido pela API, antes da extracao do JSON."),
    ("s_modelo", "Juiz secundario - modelo", 12, "glm-4.5, o proprio gerador (julgamento da execucao original)."),
    ("s_media", "Juiz secundario - media", 10, "Media simples de D1 a D5."),
    ("s_d1", "Secundario D1", 9, "Nota 1-5."),
    ("s_d2", "Secundario D2", 9, "Nota 1-5."),
    ("s_d3", "Secundario D3", 9, "Nota 1-5."),
    ("s_d4", "Secundario D4", 9, "Nota 1-5."),
    ("s_d5", "Secundario D5", 9, "Nota 1-5."),
    ("s_justificativa", "Secundario - justificativa", 40, "Justificativa textual do juiz secundario."),
    (
        "contexto",
        "Bloco [CONTEXTO] enviado ao gerador",
        50,
        "Bloco que varia entre condicoes, reconstruido com build_context_block (sem exemplos reais).",
    ),
    (
        "prompt_geracao",
        "Prompt de geracao completo",
        60,
        "Prompt integral de geracao, reconstruido com build_prompt (mesmo codigo do pipeline).",
    ),
    ("prompt_juiz", "Prompt do juiz completo", 60, "Prompt integral do LLM-as-Judge, reconstruido com build_judge_prompt."),
]
LETRA = {chave: get_column_letter(i + 1) for i, (chave, *_) in enumerate(colunas)}


def cabecalho(ws, rotulos: list[str]) -> None:
    ws.append(rotulos)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = AZUL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


def como_tabela(ws, nome: str, ncols: int, nlinhas: int) -> None:
    t = Table(displayName=nome, ref=f"A1:{get_column_letter(ncols)}{nlinhas + 1}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)


wb = Workbook()

# ============================== Leia-me ==============================
leia = wb.active
leia.title = "Leia-me"
leia.column_dimensions["A"].width = 150
textos = [
    ("Graph-Driven Exam Gen - corpus experimental auditavel (1.788 itens, 120 habilidades, 3 condicoes)", True),
    (f"Repositorio: {REPO}  |  Arquivo gerado por backend/scripts/build_xlsx.py a partir de backend/data/.", False),
    ("", False),
    ("PLANILHAS", True),
    (
        "Questoes - uma linha por item: conteudo gerado, contexto curricular, criterios C1-C5, duplicatas, metricas "
        "(gravadas e recalculadas), os dois juizes e os prompts integrais.",
        False,
    ),
    ("Resumo - contagens e medias por condicao (formulas sobre a planilha Questoes).", False),
    (
        "Habilidades - medias de C, E e do juiz principal por habilidade e condicao (base da analise pareada por habilidade).",
        False,
    ),
    (
        "Matriz - habilidades, competencias, areas e objetos de conhecimento usados no grafo; Objetos de conhecimento "
        "- textos integrais.",
        False,
    ),
    (
        "Estatistica - todos os resultados numericos dos arquivos resultado_*.json (testes, efeitos, robustez, "
        "juizes, duplicatas).",
        False,
    ),
    (
        "Prompts - templates de geracao (por condicao) e do juiz; Dicionario - descricao de cada coluna da planilha Questoes.",
        False,
    ),
    ("", False),
    ("PROVENIENCIA DE CADA BLOCO DE COLUNAS", True),
    (
        "Gravado no experimento (questoes_export.json): conteudo do item, validacao registrada, C, E, TTR, FKGL e "
        "notas/justificativa do juiz secundario.",
        False,
    ),
    (
        "Gravado no rejulgamento (julgamento_glm-4-plus.jsonl): notas, justificativa, resposta bruta, tokens, "
        "data/hora e parametros do juiz principal.",
        False,
    ),
    (
        "Reconstruido agora com o mesmo codigo do pipeline: criterios C1-C5, C e E recalculados, bloco [CONTEXTO] e "
        "prompts de geracao e do juiz.",
        False,
    ),
    (
        "Calculado agora: marcacao de enunciados repetidos e similaridade maxima dentro da celula (habilidade x condicao).",
        False,
    ),
    ("", False),
    ("O QUE NAO ESTA DISPONIVEL (e por que)", True),
    (
        "Respostas brutas, tokens e horarios da GERACAO e do JUIZ SECUNDARIO: os registros dessas etapas nao foram "
        "preservados; o periodo de geracao (05 a 10/09/2026) foi inferido das datas dos arquivos.",
        False,
    ),
    (
        "Exemplos few-shot: o bloco 'Exemplos:' foi enviado vazio em com_grafo e com_fewshot_sem_grafo, porque o "
        "corpus de itens oficiais do ENEM nao foi transcrito.",
        False,
    ),
    ("Similaridade semantica: nao calculada, pois a chave de API nao tinha acesso a modelos de embedding.", False),
    (
        "Os 12 itens faltantes (1.800 planejados) nao chegaram a ser gravados; sem os registros, a causa nao pode ser "
        "confirmada.",
        False,
    ),
    ("", False),
    ("COMO REPRODUZIR", True),
    (
        "Analises: uv run python -m scripts.analisar_corpus | scripts.analise_robustez | "
        "scripts.analisar_juiz_distinto | scripts.analise_duplicatas (a partir de backend/).",
        False,
    ),
    (
        "Esta planilha: uv run --with openpyxl python -m scripts.build_xlsx. As colunas 'recalculado' devem coincidir "
        "com as gravadas (ver planilha Estatistica, verificacao).",
        False,
    ),
]
for texto, negrito in textos:
    leia.append([texto])
    leia.cell(row=leia.max_row, column=1).font = Font(bold=negrito, size=14 if leia.max_row == 1 else 11)
    leia.cell(row=leia.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")

# ============================== Questoes ==============================
ws = wb.create_sheet("Questoes")
cabecalho(ws, [rotulo for _, rotulo, *_ in colunas])
sem_principal = 0
dif_c = dif_e = 0.0
reprovados_agora = 0
for r in registros:
    hab = habilidades[r["habilidade_codigo"]]
    cond = CondicaoExperimental(r["origem"])
    item = item_de(r)
    reprovados = set(validador.validar(item).criterios_reprovados)
    reprovados_agora += bool(reprovados)

    tokens_item = nlp.tokenizar(r["enunciado"] + " " + " ".join(r[f"alternativa_{x}"] for x in "abcde"))
    c_alvo = nlp.cobertura_lexical(tokens_item, tokens_oc[hab.habilidade_codigo])
    c_outras = [nlp.cobertura_lexical(tokens_item, t) for k, t in tokens_oc.items() if k != hab.habilidade_codigo]
    e_recalc = nlp.especificidade_curricular(c_alvo, c_outras)
    dif_c = max(dif_c, abs(c_alvo - numero(r["cobertura_lexical"])))
    dif_e = max(dif_e, abs(e_recalc - numero(r["especificidade_curricular"])))

    p = juiz_principal.get(r["id"])
    sem_principal += p is None
    uso = ast.literal_eval(p["uso_tokens"]) if p and isinstance(p.get("uso_tokens"), str) else (p or {}).get("uso_tokens") or {}
    dims_p = [numero(p.get(f"d{k}")) for k in range(1, 6)] if p else [None] * 5
    dims_s = [numero(r.get(k)) for k in DIMS_SECUNDARIO]
    extra = {
        "competencia_descricao": hab.competencia_descricao,
        "ocs": "; ".join(o.titulo for o in hab.objetos_conhecimento),
        **{chave: "Nao (reprovado)" if chave in reprovados else "Sim" for chave, _ in CRITERIOS},
        "dup": "Sim" if primeira_ocorrencia[r["id"]] else "Nao",
        "dup_de": primeira_ocorrencia[r["id"]],
        "sim_max": round(maior_similaridade[r["id"]], 4),
        "c_recalc": c_alvo,
        "e_recalc": e_recalc,
        "p_modelo": p.get("modelo_juiz") if p else None,
        "p_julgado_em": p.get("julgado_em") if p else None,
        "p_temperatura": numero(p.get("temperatura")) if p else None,
        "p_max_tokens": numero(p.get("max_tokens")) if p else None,
        "p_media": media(dims_p) if p else None,
        **{f"p_d{k}": dims_p[k - 1] for k in range(1, 6)},
        "p_justificativa": p.get("justificativa") if p else None,
        "p_tok_prompt": uso.get("prompt_tokens"),
        "p_tok_resp": uso.get("completion_tokens"),
        "p_tok_total": uso.get("total_tokens"),
        "p_bruta": p.get("resposta_bruta") if p else None,
        "s_modelo": "glm-4.5",
        "s_media": media(dims_s),
        **{f"s_d{k}": dims_s[k - 1] for k in range(1, 6)},
        "s_justificativa": r.get("judge_justificativa"),
        "contexto": build_context_block(cond, hab, []),
        "prompt_geracao": build_prompt(cond, hab, []),
        "prompt_juiz": build_judge_prompt(item, hab),
    }
    linha = []
    for chave, *_ in colunas:
        v = extra[chave] if chave in extra else r.get(chave)
        if chave == "validacao_estrutural_ok":
            v = "Sim" if v in (True, "True", "true", 1) else "Nao"
        elif chave in ("competencia_numero", "cobertura_lexical", "especificidade_curricular", "ttr", "fkgl"):
            v = numero(v)
        linha.append(v)
    ws.append(linha)

for chave, _, largura, _ in colunas:
    ws.column_dimensions[LETRA[chave]].width = largura
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=False)
como_tabela(ws, "TabelaQuestoes", len(colunas), len(registros))

# ============================== Resumo ==============================
resumo = wb.create_sheet("Resumo", 1)
resumo["A1"] = "Resumo por condicao (formulas sobre a planilha Questoes)"
resumo["A1"].font = Font(bold=True, size=14)
rotulos = [
    "Condicao",
    "Qtd. itens",
    "Enunciados repetidos",
    "C(i,h) media",
    "E(i,h) media",
    "TTR media",
    "FKGL media",
    "Juiz principal (glm-4-plus) media",
    "Juiz secundario (glm-4.5) media",
]
chaves = [None, None, None, "cobertura_lexical", "especificidade_curricular", "ttr", "fkgl", "p_media", "s_media"]
for j, texto in enumerate(rotulos):
    c = resumo.cell(row=3, column=j + 1, value=texto)
    c.font = Font(bold=True)
    c.alignment = Alignment(wrap_text=True)
    resumo.column_dimensions[get_column_letter(j + 1)].width = 24 if j == 0 else 16
cc, dc = LETRA["origem"], LETRA["dup"]
for i, cond in enumerate([*CONDICOES, "Total"]):
    row = 4 + i
    resumo.cell(row=row, column=1, value=cond)
    if cond == "Total":
        resumo.cell(row=row, column=2, value=f"=COUNTA(Questoes!{LETRA['id']}2:{LETRA['id']}100000)")
        resumo.cell(row=row, column=3, value=f'=COUNTIF(Questoes!{dc}2:{dc}100000,"Sim")')
        for j, chave in enumerate(chaves[3:], start=4):
            col = LETRA[chave]
            resumo.cell(row=row, column=j, value=f"=AVERAGE(Questoes!{col}2:{col}100000)")
        continue
    resumo.cell(row=row, column=2, value=f'=COUNTIF(Questoes!{cc}2:{cc}100000,"{cond}")')
    resumo.cell(row=row, column=3, value=f'=COUNTIFS(Questoes!{cc}2:{cc}100000,"{cond}",Questoes!{dc}2:{dc}100000,"Sim")')
    for j, chave in enumerate(chaves[3:], start=4):
        col = LETRA[chave]
        resumo.cell(row=row, column=j, value=f'=AVERAGEIF(Questoes!{cc}2:{cc}100000,"{cond}",Questoes!{col}2:{col}100000)')

# ============================== Habilidades ==============================
hs = wb.create_sheet("Habilidades")
rot = ["Habilidade", "Area", "Competencia", "Descricao"]
for cond in CONDICOES:
    rot += [f"{cond} - n", f"{cond} - C media", f"{cond} - E media", f"{cond} - juiz principal media"]
rot += ["C: grafo > fewshot", "C: grafo > sem_grafo", "E: grafo > fewshot", "E: grafo > sem_grafo"]
cabecalho(hs, rot)


def med(vals):
    vals = [numero(v) for v in vals if v is not None]
    return sum(vals) / len(vals) if vals else None


for codigo, hab in habilidades.items():
    linha = [codigo, hab.area_codigo, hab.competencia_numero, hab.habilidade_descricao]
    medias = {}
    for cond in CONDICOES:
        itens = por_celula.get((codigo, cond), [])
        mc = med([q["cobertura_lexical"] for q in itens])
        me = med([q["especificidade_curricular"] for q in itens])
        mj = med(
            [media([juiz_principal[q["id"]].get(f"d{k}") for k in range(1, 6)]) for q in itens if q["id"] in juiz_principal]
        )
        medias[cond] = (mc, me)
        linha += [len(itens), mc, me, mj]
    for idx in (0, 1):
        for outra in ("com_fewshot_sem_grafo", "sem_grafo"):
            a, b = medias["com_grafo"][idx], medias[outra][idx]
            linha.append(None if a is None or b is None else ("Sim" if a > b else "Nao"))
    hs.append(linha)
for j in range(1, len(rot) + 1):
    hs.column_dimensions[get_column_letter(j)].width = 40 if j == 4 else 14
como_tabela(hs, "TabelaHabilidades", len(rot), len(habilidades))

# ============================== Matriz e Objetos de conhecimento ==============================
ms = wb.create_sheet("Matriz")
cabecalho(
    ms,
    [
        "Area",
        "Area (nome)",
        "Competencia",
        "Competencia (descricao)",
        "Habilidade",
        "Habilidade (descricao)",
        "Objetos de conhecimento (codigos)",
    ],
)
for area in matriz["areas"]:
    for comp in area["competencias"]:
        for hab in comp["habilidades"]:
            ms.append(
                [
                    area["codigo"],
                    area["nome"],
                    comp["numero"],
                    comp["descricao"],
                    hab["codigo"],
                    hab["descricao"],
                    "; ".join(comp["objetos_conhecimento"]),
                ]
            )
for j, w in enumerate([8, 24, 11, 60, 11, 60, 30], start=1):
    ms.column_dimensions[get_column_letter(j)].width = w
como_tabela(ms, "TabelaMatriz", 7, ms.max_row - 1)

os_ = wb.create_sheet("Objetos de conhecimento")
cabecalho(os_, ["Area", "Codigo", "Titulo", "Descricao (como usada no grafo e em O(h))"])
for area in matriz["areas"]:
    for oc in area["objetos_conhecimento"]:
        os_.append([area["codigo"], oc["codigo"], oc["titulo"], oc["descricao"]])
for j, w in enumerate([8, 12, 40, 120], start=1):
    os_.column_dimensions[get_column_letter(j)].width = w
como_tabela(os_, "TabelaObjetos", 4, os_.max_row - 1)

# ============================== Estatistica ==============================
es = wb.create_sheet("Estatistica")
cabecalho(es, ["Arquivo de origem", "Caminho do resultado", "Valor"])


def achatar(prefixo: str, valor, arquivo: str) -> None:
    if isinstance(valor, dict):
        for k, v in valor.items():
            achatar(f"{prefixo}.{k}" if prefixo else str(k), v, arquivo)
    elif isinstance(valor, list) and valor and all(isinstance(v, dict) for v in valor):
        for i, v in enumerate(valor):
            achatar(f"{prefixo}[{i}]", v, arquivo)
    else:
        es.append([arquivo, prefixo, json.dumps(valor, ensure_ascii=False) if isinstance(valor, list) else valor])


for nome in [
    "resultado_analise_estatistica.json",
    "resultado_robustez.json",
    "resultado_juiz_distinto_glm-4-plus.json",
    "resultado_robustez_duplicatas.json",
    "conferencia_extracao_matriz.json",
]:
    if (DATA_DIR / nome).exists():
        achatar("", ler_json(nome), nome)
es.append(["verificacao (build_xlsx.py)", "max |C gravado - C recalculado|", dif_c])
es.append(["verificacao (build_xlsx.py)", "max |E gravado - E recalculado|", dif_e])
es.append(["verificacao (build_xlsx.py)", "itens reprovados ao reaplicar C1-C5", reprovados_agora])
es.append(["verificacao (build_xlsx.py)", "itens sem julgamento do juiz principal", sem_principal])
for j, w in enumerate([38, 110, 30], start=1):
    es.column_dimensions[get_column_letter(j)].width = w
como_tabela(es, "TabelaEstatistica", 3, es.max_row - 1)

# ============================== Prompts ==============================
ps = wb.create_sheet("Prompts")
cabecalho(ps, ["Uso", "Template (exemplo com CN-H19)"])
exemplo = habilidades["CN-H19"]
for cond in CONDICOES:
    ps.append([f"Geracao - {cond}", build_prompt(CondicaoExperimental(cond), exemplo, [])])
ps.append(
    [
        "Juiz (LLM-as-Judge)",
        build_judge_prompt(
            ItemGerado(
                enunciado="<enunciado>",
                alternativas={x: f"<alternativa {x}>" for x in "ABCDE"},
                gabarito="A",
                justificativa="<justificativa do gerador>",
            ),
            exemplo,
        ),
    ]
)
ps.column_dimensions["A"].width = 28
ps.column_dimensions["B"].width = 140
for row in ps.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical="top")

# ============================== Dicionario ==============================
ds = wb.create_sheet("Dicionario")
cabecalho(ds, ["Coluna (planilha Questoes)", "Descricao"])
for _, rotulo, _, descricao in colunas:
    ds.append([rotulo, descricao])
ds.column_dimensions["A"].width = 45
ds.column_dimensions["B"].width = 120

out_path = DATA_DIR / "questoes_geradas.xlsx"
wb.save(out_path)
print(
    f"Salvo em {out_path} ({len(registros)} linhas; sem juiz principal: {sem_principal}; "
    f"max dif C={dif_c:.2e}, E={dif_e:.2e}; reprovados ao reaplicar C1-C5: {reprovados_agora})"
)
